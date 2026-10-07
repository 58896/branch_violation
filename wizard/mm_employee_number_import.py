# -*- coding: utf-8 -*-
"""One-time import of Jisr employee numbers into hr.employee.mm_employee_number.

Steps: upload an xlsx -> review the matches (fix unmatched / ambiguous rows
by hand, or skip them) -> apply. An xlsx report of the rows that need
attention can be downloaded at the review and done steps.

Matching, per row: an employee who already holds the row's number (so the
wizard can be re-run safely), otherwise the employee's name after Arabic
normalisation. Name matching is only used here, to seed the numbers once;
everything else identifies employees by number.
"""
import base64
import io
import re
import unicodedata

from odoo import _, api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:  # pragma: no cover - declared in external_dependencies
    openpyxl = None
try:
    import xlsxwriter
except ImportError:  # pragma: no cover
    xlsxwriter = None

HEADER_ROWS_TO_SCAN = 15
HEADER_ALIASES = {
    "number": {
        "الرقم الوظيفي", "رقم الموظف", "الرقم", "employee number",
        "employee no", "employee no.", "employee id", "emp no", "emp id",
        "employee code", "number",
    },
    "name": {
        "اسم الموظف", "الاسم", "الاسم الكامل", "employee name", "name",
        "full name",
    },
    "job": {
        "المسمى الوظيفي", "المسمي الوظيفي", "الوظيفة", "job title", "job",
        "position", "job position",
    },
    "branch": {
        "الفرع", "المعرض", "اسم المعرض", "branch", "branch name", "location",
    },
}

# Summary rows some exports add under the employee list
TOTAL_ROW_LABELS = {"الاجمالي", "المجموع", "total", "grand total"}

MATCH_STATUSES = [
    ("matched", "Matched"),
    ("manual", "Chosen Manually"),
    ("ambiguous", "Ambiguous"),
    ("unmatched", "Not Found"),
    ("skipped", "Skipped"),
]

_ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")


def normalize_text(value):
    """Loose comparison key for Arabic/English names and headers."""
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = re.sub("[ً-ْٰـ]", "", text)  # harakat, tatweel
    text = re.sub("[أإآٱ]", "ا", text)
    text = text.replace("ى", "ي").replace("ة", "ه")
    text = re.sub(r"[‐-―−]", "-", text)  # unicode dashes
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text.replace("عبد ", "عبد")


def normalize_number(value):
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).translate(_ARABIC_DIGITS).strip()


class MmEmployeeNumberImport(models.TransientModel):
    _name = "mm.employee.number.import"
    _description = "Import Employee Numbers"

    state = fields.Selection(
        [("upload", "Upload"), ("review", "Review"), ("done", "Done")],
        default="upload", required=True,
    )
    company_id = fields.Many2one(
        "res.company", required=True, default=lambda self: self.env.company,
    )
    file = fields.Binary(string="Excel File", attachment=False)
    filename = fields.Char()
    sheet_name = fields.Char(
        help="Leave empty to use the sheet with the most employee rows.",
    )
    update_job_branch = fields.Boolean(
        string="Also update job position and branch",
        help="Also set each matched employee's job position (created if it "
             "does not exist yet) and branch, when the file has those columns.",
    )
    line_ids = fields.One2many("mm.employee.number.import.line", "import_id")
    used_sheet = fields.Char(readonly=True)
    has_job_column = fields.Boolean(readonly=True)
    has_branch_column = fields.Boolean(readonly=True)

    matched_count = fields.Integer(compute="_compute_counts")
    attention_count = fields.Integer(compute="_compute_counts")
    skipped_count = fields.Integer(compute="_compute_counts")

    result_message = fields.Text(readonly=True)
    report_file = fields.Binary(readonly=True, attachment=False)
    report_filename = fields.Char()

    @api.depends("line_ids.match_status", "line_ids.issue")
    def _compute_counts(self):
        for wiz in self:
            lines = wiz.line_ids
            wiz.matched_count = len(lines.filtered(
                lambda l: l.match_status in ("matched", "manual") and not l.issue))
            wiz.skipped_count = len(lines.filtered(lambda l: l.match_status == "skipped"))
            wiz.attention_count = len(lines) - wiz.matched_count - wiz.skipped_count

    def _reopen(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
            "name": _("Import Employee Numbers"),
        }

    # ------------------------------------------------------------------
    # Step 1: read the file
    # ------------------------------------------------------------------
    def _read_rows(self):
        """Return (sheet_title, columns, rows); rows are dicts keyed by the
        HEADER_ALIASES keys plus 'row'."""
        self.ensure_one()
        if not openpyxl:
            raise UserError(_("The openpyxl Python library is not installed."))
        if not self.file:
            raise UserError(_("Upload an Excel file first."))
        try:
            book = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.file)), read_only=True, data_only=True)
        except Exception as exc:
            raise UserError(_("Could not read the file as .xlsx: %s", exc)) from exc
        aliases = {key: {normalize_text(a) for a in values}
                   for key, values in HEADER_ALIASES.items()}
        sheets = book.worksheets
        if self.sheet_name:
            sheets = [s for s in sheets if s.title.strip() == self.sheet_name.strip()]
            if not sheets:
                raise UserError(_("The file has no sheet named “%s”.", self.sheet_name))
        best = None
        for sheet in sheets:
            all_rows = list(sheet.iter_rows(values_only=True))
            for header_idx, header in enumerate(all_rows[:HEADER_ROWS_TO_SCAN]):
                columns = {}
                for col_idx, cell in enumerate(header):
                    key = normalize_text(cell)
                    for field_key, names in aliases.items():
                        if key in names and field_key not in columns:
                            columns[field_key] = col_idx
                if "number" not in columns or "name" not in columns:
                    continue
                rows = []
                for offset, values in enumerate(all_rows[header_idx + 1:], start=header_idx + 2):
                    get = lambda k: values[columns[k]] if k in columns and columns[k] < len(values) else None
                    number, name = normalize_number(get("number")), str(get("name") or "").strip()
                    if not (number or name) or number.startswith("=") or \
                            (not number and normalize_text(name) in TOTAL_ROW_LABELS):
                        continue
                    rows.append({
                        "row": offset, "number": number, "name": name,
                        "job": str(get("job") or "").strip(),
                        "branch": str(get("branch") or "").strip(),
                    })
                if best is None or len(rows) > len(best[2]):
                    best = (sheet.title, columns, rows)
                break
        if not best or not best[2]:
            raise UserError(_(
                "No employee rows found. The sheet needs a header row with an "
                "employee number column (e.g. “الرقم الوظيفي”) and an employee "
                "name column (e.g. “اسم الموظف”)."))
        return best

    def action_read_file(self):
        self.ensure_one()
        self._mm_check_access()
        title, columns, rows = self._read_rows()
        employees = self.env["hr.employee"].sudo().with_context(active_test=False).search(
            [("company_id", "=", self.company_id.id)])
        by_number, by_name = {}, {}
        for emp in employees:
            if emp.mm_employee_number:
                by_number[emp.mm_employee_number] = emp
            by_name.setdefault(normalize_text(emp.name), self.env["hr.employee"])
            by_name[normalize_text(emp.name)] |= emp
        branches = self._mm_branch_index()

        vals_list = []
        for row in rows:
            employee = by_number.get(row["number"]) if row["number"] else None
            candidates = self.env["hr.employee"]
            if not employee:
                candidates = by_name.get(normalize_text(row["name"]), self.env["hr.employee"])
                active = candidates.filtered("active")
                if len(active) == 1:
                    employee = active
                elif not active and len(candidates) == 1:
                    employee = candidates
            vals_list.append({
                "import_id": self.id,
                "row_number": row["row"],
                "file_employee_number": row["number"],
                "file_employee_name": row["name"],
                "file_job": row["job"],
                "file_branch": row["branch"],
                "auto_employee_id": employee.id if employee else False,
                "employee_id": employee.id if employee else False,
                "candidate_ids": [(6, 0, candidates.ids if not employee and len(candidates) > 1 else [])],
                "branch_id": branches.get(self._mm_branch_key(row["branch"])) if row["branch"] else False,
            })
        self.line_ids.unlink()
        self.env["mm.employee.number.import.line"].create(vals_list)
        self.write({
            "state": "review",
            "used_sheet": title,
            "has_job_column": "job" in columns,
            "has_branch_column": "branch" in columns,
        })
        self._mm_generate_report()
        return self._reopen()

    @api.model
    def _mm_branch_key(self, value):
        value = str(value or "").strip()
        bracket = re.match(r"^\[(.+?)\]", value)
        return normalize_text(bracket.group(1) if bracket else value)

    def _mm_branch_index(self):
        """normalised code / name (every installed language) -> branch id."""
        index = {}
        Branch = self.env["company.branch"].sudo().with_context(active_test=False)
        branches = Branch.search([("company_id", "=", self.company_id.id)])
        for lang, _name in self.env["res.lang"].get_installed():
            for branch in branches.with_context(lang=lang):
                index.setdefault(normalize_text(branch.name), branch.id)
        for branch in branches:
            index[normalize_text(branch.code)] = branch.id
        return index

    # ------------------------------------------------------------------
    # Step 2: review + apply
    # ------------------------------------------------------------------
    def action_back(self):
        self.ensure_one()
        self.line_ids.unlink()
        self.write({"state": "upload", "report_file": False, "result_message": False})
        return self._reopen()

    def action_refresh_report(self):
        self.ensure_one()
        self._mm_generate_report()
        return self._reopen()

    def action_apply(self):
        self.ensure_one()
        self._mm_check_access()
        lines = self.line_ids.filtered(lambda l: l.match_status in ("matched", "manual"))
        blocking = lines.filtered("issue")
        if blocking:
            rows = [self.env._("Row %(row)s (%(name)s): %(issue)s",
                               row=l.row_number, name=l.file_employee_name, issue=l.issue)
                    for l in blocking[:20]]
            raise UserError(_("Fix or skip these rows first:\n%s", "\n".join(rows)))
        if not lines:
            raise UserError(_("There are no matched rows to apply."))

        Employee = self.env["hr.employee"].sudo()
        changed = lines.filtered(
            lambda l: l.employee_id.mm_employee_number != l.file_employee_number)
        # Clear first so numbers can move between employees without tripping
        # the per-company unique constraint mid-way.
        changed.employee_id.sudo().write({"mm_employee_number": False})
        for line in changed:
            line.employee_id.sudo().write({"mm_employee_number": line.file_employee_number})

        notes = []
        if self.update_job_branch:
            notes += self._mm_apply_job_branch(lines)

        self.write({
            "state": "done",
            "result_message": "\n".join([
                _("Employee numbers set or changed: %s", len(changed)),
                _("Already up to date: %s", len(lines) - len(changed)),
                _("Rows not imported (need attention or skipped): %s",
                  len(self.line_ids) - len(lines)),
            ] + notes),
        })
        self._mm_generate_report()
        Employee.flush_model()
        return self._reopen()

    def _mm_apply_job_branch(self, lines):
        notes = []
        Job = self.env["hr.job"].sudo()
        created_jobs = Job
        job_cache = {}
        for line in lines.filtered("file_job"):
            key = normalize_text(line.file_job)
            if key not in job_cache:
                job = Job.search([
                    ("company_id", "in", [False, self.company_id.id]),
                ]).filtered(lambda j: normalize_text(j.name) == key)[:1]
                if not job:
                    job = Job.create({"name": line.file_job, "company_id": self.company_id.id})
                    created_jobs |= job
                job_cache[key] = job
            line.employee_id.sudo().job_id = job_cache[key]
        if created_jobs:
            notes.append(_("Job positions created: %s", ", ".join(created_jobs.mapped("name"))))
        notes.append(_("Job positions updated: %s", len(lines.filtered("file_job"))))

        branch_field = self.env["hr.employee"]._fields.get("branch_id")
        branch_lines = lines.filtered("branch_id")
        if branch_field is not None and branch_field.comodel_name == "company.branch":
            for line in branch_lines:
                line.employee_id.sudo().branch_id = line.branch_id
            notes.append(_("Branches updated: %s", len(branch_lines)))
        elif branch_lines:
            notes.append(_(
                "Branches were NOT saved: employees have no Branch field "
                "(it comes with the POS Sales Target & Commission module)."))
        return notes

    def _mm_check_access(self):
        if not self.env.user.has_group("mm_branch_violation.group_violation_manager"):
            raise UserError(_("Only a Violation Manager can import employee numbers."))

    # ------------------------------------------------------------------
    # Report
    # ------------------------------------------------------------------
    def _mm_generate_report(self):
        self.ensure_one()
        if not xlsxwriter:
            return
        lines = self.line_ids.filtered(
            lambda l: l.match_status not in ("matched", "manual") or l.issue)
        status_labels = dict(self.env["mm.employee.number.import.line"]._fields[
            "match_status"]._description_selection(self.env))
        buf = io.BytesIO()
        book = xlsxwriter.Workbook(buf, {"in_memory": True})
        sheet = book.add_worksheet(_("Needs Attention")[:31])
        if (self.env.user.lang or "").startswith("ar"):
            sheet.right_to_left()
        bold = book.add_format({"bold": True, "bg_color": "#DDDDDD", "border": 1})
        headers = [
            _("Row"), _("Employee Number"), _("Name in File"), _("Job in File"),
            _("Branch in File"), _("Status"), _("Issue"), _("Possible Employees"),
        ]
        for col, head in enumerate(headers):
            sheet.write(0, col, head, bold)
        for row, line in enumerate(lines, start=1):
            sheet.write_number(row, 0, line.row_number)
            sheet.write_string(row, 1, line.file_employee_number or "")
            sheet.write_string(row, 2, line.file_employee_name or "")
            sheet.write_string(row, 3, line.file_job or "")
            sheet.write_string(row, 4, line.file_branch or "")
            sheet.write_string(row, 5, status_labels.get(line.match_status, ""))
            sheet.write_string(row, 6, line.issue or "")
            sheet.write_string(row, 7, ", ".join(
                "%s (#%s)" % (e.name, e.id) for e in line.sudo().candidate_ids))
        sheet.set_column(0, 0, 6)
        sheet.set_column(1, 1, 16)
        sheet.set_column(2, 7, 32)
        sheet.freeze_panes(1, 0)
        book.close()
        self.write({
            "report_file": base64.b64encode(buf.getvalue()),
            "report_filename": "employee_number_import_report.xlsx",
        })


class MmEmployeeNumberImportLine(models.TransientModel):
    _name = "mm.employee.number.import.line"
    _description = "Import Employee Numbers - Row"
    _order = "row_number"

    import_id = fields.Many2one(
        "mm.employee.number.import", required=True, ondelete="cascade",
    )
    company_id = fields.Many2one(related="import_id.company_id")
    row_number = fields.Integer(string="Row", readonly=True)
    file_employee_number = fields.Char(string="Employee Number", readonly=True)
    file_employee_name = fields.Char(string="Name in File", readonly=True)
    file_job = fields.Char(string="Job in File", readonly=True)
    file_branch = fields.Char(string="Branch in File", readonly=True)
    auto_employee_id = fields.Many2one("hr.employee", readonly=True)
    candidate_ids = fields.Many2many(
        "hr.employee", string="Possible Employees", readonly=True,
        context={"active_test": False},
    )
    employee_id = fields.Many2one(
        "hr.employee", string="Odoo Employee",
        domain="[('company_id', '=', company_id)]",
        context={"active_test": False},
    )
    employee_active = fields.Boolean(related="employee_id.active")
    branch_id = fields.Many2one(
        "company.branch", string="Branch",
        domain="[('company_id', '=', company_id)]",
    )
    skip = fields.Boolean()
    match_status = fields.Selection(
        MATCH_STATUSES, compute="_compute_match_status", store=True,
    )
    issue = fields.Char(compute="_compute_issue")

    @api.depends("skip", "employee_id", "auto_employee_id", "candidate_ids")
    def _compute_match_status(self):
        for line in self:
            if line.skip:
                line.match_status = "skipped"
            elif line.employee_id:
                line.match_status = (
                    "matched" if line.employee_id == line.auto_employee_id else "manual")
            elif line.candidate_ids:
                line.match_status = "ambiguous"
            else:
                line.match_status = "unmatched"

    @api.depends("skip", "employee_id", "file_employee_number", "file_branch",
                 "branch_id", "import_id.line_ids.employee_id",
                 "import_id.line_ids.skip", "import_id.update_job_branch")

    def _compute_issue(self):
        imports = self.import_id
        holders = {}
        numbers = [n for n in imports.line_ids.mapped("file_employee_number") if n]
        if numbers:
            for emp in self.env["hr.employee"].sudo().with_context(active_test=False).search([
                    ("mm_employee_number", "in", numbers),
                    ("company_id", "in", imports.company_id.ids)]):
                holders[(emp.company_id.id, emp.mm_employee_number)] = emp
        for line in self:
            issues = []
            siblings = line.import_id.line_ids.filtered(lambda l: l != line and not l.skip)
            if line.skip:
                line.issue = False
                continue
            if not line.file_employee_number:
                issues.append(_("No employee number in the file."))
            elif line.file_employee_number in siblings.mapped("file_employee_number"):
                issues.append(_("The same employee number appears on another row."))
            if line.employee_id:
                if line.employee_id in siblings.employee_id:
                    issues.append(_("The same employee is chosen on another row."))
                holder = holders.get((line.company_id.id, line.file_employee_number))
                if holder and holder != line.employee_id and \
                        holder not in siblings.employee_id:
                    issues.append(_("Number already belongs to %s.", holder.name))
            if line.import_id.update_job_branch and line.file_branch and not line.branch_id:
                issues.append(_("Branch not found - pick it or skip the row."))
            line.issue = " ".join(issues) or False
