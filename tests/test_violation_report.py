# -*- coding: utf-8 -*-
import base64
import io
from datetime import date

import openpyxl

from odoo.exceptions import UserError
from odoo.tests import tagged

from .common import MmViolationCommon


@tagged("post_install", "-at_install")
class TestViolationReport(MmViolationCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.branch2 = cls.env["company.branch"].create({
            "name": "MMV Second Branch", "code": "MMV-2ND", "company_id": cls.company.id})
        cls.employee.mm_employee_number = "R-1"
        cls.seller = cls.env["hr.employee"].create({"name": "MMV Report Seller", "mm_employee_number": "R-2"})

    def _approved(self, day, employee=None, branch=None, vtype=None, points=None):
        vals = {"date": day, "violation_type_id": (vtype or self.type_5).id}
        if employee:
            vals["employee_id"] = employee.id
        if branch:
            vals["branch_id"] = branch.id
        if points is not None:
            violation = self._create(user=self.manager, points=points, **vals)
            violation.with_user(self.manager).action_submit()
            violation.with_user(self.manager2).action_approve()
        else:
            violation = self._create(user=self.officer, **vals)
            violation.with_user(self.officer).action_submit()
            violation.with_user(self.manager).action_approve()
        return violation

    def _wizard(self, **vals):
        return self.env["mm.violation.report.wizard"].with_user(self.officer).create(dict({
            "period": "custom", "date_from": date(2032, 3, 1), "date_to": date(2032, 4, 30),
            "language": "en_US",
        }, **vals))

    def _data(self):
        self._approved(date(2032, 3, 5), vtype=self.type_10)
        self._approved(date(2032, 3, 9))
        self._approved(date(2032, 4, 2), employee=self.seller, branch=self.branch2, points=40)
        self._create(user=self.officer, date=date(2032, 3, 10))   # draft: pending only

    def test_sections(self):
        self._data()
        data = self._wizard()._mm_collect()
        kpis = data["kpis"]
        self.assertEqual(kpis["violations"], 3)
        self.assertEqual(kpis["points"], 55.0)
        self.assertEqual(kpis["employees"], 2)
        self.assertEqual(kpis["branches"], 2)
        self.assertEqual(kpis["critical"], 1)
        # ranking: seller (40) first, critical; cashier (15) warning
        first, second = data["employees"]
        self.assertEqual((first["number"], first["rank"], first["level"]), ("R-2", 1, "critical"))
        self.assertEqual((second["points"], second["level"]), (15.0, "warning"))
        self.assertEqual(data["branches"][0]["name"], self.branch2.display_name)
        self.assertEqual(sum(r["violations"] for r in data["types"]), 3)
        self.assertEqual([r["month"] for r in data["trend"]], ["2032-03", "2032-04"])
        self.assertTrue(data["show_trend"])
        self.assertEqual(len(data["details"]), 3)

    def test_filters(self):
        self._data()
        data = self._wizard(branch_ids=[(6, 0, self.branch.ids)])._mm_collect()
        self.assertEqual(data["kpis"]["violations"], 2)
        data = self._wizard(violation_type_ids=[(6, 0, self.type_10.ids)])._mm_collect()
        self.assertEqual(data["kpis"]["points"], 10.0)
        data = self._wizard(status_filter="pending")._mm_collect()
        self.assertEqual(data["kpis"]["violations"], 1)
        data = self._wizard(include_details=False)._mm_collect()
        self.assertFalse(data["details"])
        self.assertIn(self.branch.display_name, self._wizard(
            branch_ids=[(6, 0, self.branch.ids)])._mm_filters_label())

    def test_period_presets(self):
        wizard = self.env["mm.violation.report.wizard"].create({"period": "last_month"})
        self.assertEqual(wizard.date_from.day, 1)
        self.assertLess(wizard.date_to, wizard.date_from.replace(day=28) + __import__("datetime").timedelta(days=5))
        with self.assertRaises(UserError):
            self._wizard(date_from=date(2032, 5, 1), date_to=date(2032, 4, 1))._mm_domain()

    def test_pdf_and_excel(self):
        self._data()
        wizard = self._wizard()
        # test mode renders the report as HTML instead of calling wkhtmltopdf;
        # same language context as action_print_pdf
        html, _fmt = self.env["ir.actions.report"].with_context(lang=wizard.language)._render_qweb_pdf(
            "mm_branch_violation.action_report_mm_violation", wizard.ids)
        html = html.decode()
        for text in ("Branch Violations Report", "Employee Ranking", "R-2", "Critical", "Monthly Trend"):
            self.assertIn(text, html)
        action = wizard.action_export_xlsx()
        attachment = self.env["ir.attachment"].browse(int(action["url"].split("/")[3].split("?")[0]))
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(attachment.datas)))
        self.assertEqual(book.sheetnames, ["Summary", "chart_data", "By Branch", "By Violation", "By Employee",
                                           "By Source", "Monthly Trend", "Details"])
        ranking = list(book["By Employee"].iter_rows(min_row=5, max_row=6, values_only=True))
        self.assertEqual(ranking[0][1:3], ("R-2", "MMV Report Seller"))
        self.assertEqual(ranking[0][8], "Critical")

    def test_empty_report(self):
        wizard = self._wizard(date_from=date(2040, 1, 1), date_to=date(2040, 1, 31))
        self.assertEqual(wizard._mm_collect()["kpis"]["violations"], 0)
        html, _fmt = self.env["ir.actions.report"].with_context(lang=wizard.language)._render_qweb_pdf(
            "mm_branch_violation.action_report_mm_violation", wizard.ids)
        self.assertIn("No violations match", html.decode())
        wizard.action_export_xlsx()   # must not crash on empty sections
